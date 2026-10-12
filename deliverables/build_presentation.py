from pathlib import Path
from math import floor

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REV = Path(r"D:\workspace\repositories\reverse-registration")
OUT = ROOT / "deliverables"
OUT.mkdir(exist_ok=True)

W, H = 13.333, 7.5
NAVY = RGBColor(10, 22, 40)
NAVY2 = RGBColor(16, 34, 56)
TEAL = RGBColor(35, 211, 176)
CYAN = RGBColor(59, 180, 255)
WHITE = RGBColor(245, 248, 252)
MUTED = RGBColor(169, 185, 202)
LINE = RGBColor(44, 66, 88)
YELLOW = RGBColor(255, 199, 92)
RED = RGBColor(255, 111, 116)
GREEN = RGBColor(112, 222, 142)
FONT = "Malgun Gothic"
MONO = "Consolas"


def set_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, fill=NAVY2, line=LINE, radius=True, transparency=0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def text(slide, txt, x, y, w, h, size=18, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, margin=0.05, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = txt
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def rich_lines(slide, lines, x, y, w, h, size=16, gap=7, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, item in enumerate(lines):
        if len(item) == 2:
            lead, body = item
            color = WHITE
        else:
            lead, body, color = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        p.text = ("• " if bullet else "")
        if p.text:
            p.runs[0].font.color.rgb = TEAL
            p.runs[0].font.size = Pt(size)
        r1 = p.add_run(); r1.text = lead
        r1.font.name = FONT; r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = color
        r2 = p.add_run(); r2.text = body
        r2.font.name = FONT; r2.font.size = Pt(size); r2.font.color.rgb = MUTED
    return box


def title(slide, kicker, headline, page):
    text(slide, kicker.upper(), 0.6, 0.28, 8.5, 0.28, 10, TEAL, True)
    text(slide, headline, 0.6, 0.62, 11.9, 0.62, 27, WHITE, True)
    text(slide, f"0{page}", 12.25, 0.28, 0.48, 0.3, 10, MUTED, True, PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.28), Inches(12.15), Inches(0.015))
    line.fill.solid(); line.fill.fore_color.rgb = LINE; line.line.fill.background()


def footer(slide, source):
    text(slide, source, 0.62, 7.18, 11.8, 0.18, 8.5, MUTED)


def add_image_cover(slide, path, x, y, w, h, crop=True):
    path = Path(path)
    if not path.exists():
        rect(slide, x, y, w, h)
        text(slide, f"missing: {path.name}", x, y+h/2-.15, w, .3, 10, MUTED, align=PP_ALIGN.CENTER)
        return
    if not crop:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        return
    with Image.open(path) as im:
        iw, ih = im.size
    ir, br = iw / ih, w / h
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if ir > br:
        visible = br / ir
        pic.crop_left = pic.crop_right = (1 - visible) / 2
    else:
        visible = ir / br
        pic.crop_top = pic.crop_bottom = (1 - visible) / 2


def pill(slide, txt, x, y, w, color=TEAL, txt_color=NAVY):
    s = rect(slide, x, y, w, .34, color, color)
    text(slide, txt, x, y+.02, w, .22, 10, txt_color, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return s


def arrow(slide, x, y, w, color=TEAL):
    s = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(.38))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()


def codebox(slide, code, x, y, w, h, accent=TEAL):
    rect(slide, x, y, w, h, RGBColor(7, 15, 27), LINE)
    text(slide, code, x+.15, y+.14, w-.3, h-.25, 10.5, RGBColor(205, 222, 236), False, font=MONO)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(.045), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
blank = prs.slide_layouts[6]

# 1 — X-ray numbering overview
s = prs.slides.add_slide(blank); set_bg(s); title(s, "X-RAY TEETH NUMBERING · RF-DETR", "파노라마 X-ray에서 치아를 ‘검출’하고 ‘번호화’한다", 1)
rect(s, .58, 1.52, 7.35, 4.46, RGBColor(7, 15, 27), LINE)
add_image_cover(s, ROOT / "sample.jpg", .72, 1.66, 7.07, 3.35)
pill(s, "INPUT · PANORAMIC X-RAY", .9, 1.83, 2.25, CYAN, NAVY)
text(s, "치아별 box + class + optional mask", .86, 5.18, 3.5, .28, 14, WHITE, True)
text(s, "출력 class는 모델 index → FDI 번호로 재매핑", .86, 5.53, 6.5, .25, 11, MUTED)

rect(s, 8.18, 1.52, 4.57, 1.25)
text(s, "WHY", 8.43, 1.72, .7, .22, 10, TEAL, True)
text(s, "치아 위치·종류·결손을 한 번에 구조화", 8.43, 2.02, 3.95, .38, 18, WHITE, True)
rect(s, 8.18, 2.94, 4.57, 1.50)
text(s, "FDI 32-TOOTH SYSTEM", 8.43, 3.15, 2.2, .22, 10, TEAL, True)
text(s, "상악 18→11 | 21→28", 8.43, 3.50, 3.8, .28, 16, WHITE, True)
text(s, "하악 48→41 | 31→38", 8.43, 3.87, 3.8, .28, 16, WHITE, True)
rect(s, 8.18, 4.62, 4.57, 1.36)
rich_lines(s, [("33-class 구성", " = background/기타 + 32개 치아"),
               ("두 가지 출력", " detection / instance mask")], 8.35, 4.82, 4.08, .9, 13, 5)

for i, (lab, col) in enumerate([("DATA", CYAN), ("RF-DETR", TEAL), ("FDI MAP", YELLOW), ("RESULT", GREEN)]):
    x = .82 + i*2.9
    pill(s, lab, x, 6.27, 1.65, col, NAVY)
    if i < 3: arrow(s, x+1.83, 6.25, .65, MUTED)
footer(s, "근거: rfdetr/datasets/xraypanoramic.py · rfdetr/util/teeth_classes.py · branch detr_teeth")

# 2 — implementation + lightweight
s = prs.slides.add_slide(blank); set_bg(s); title(s, "X-RAY TEETH NUMBERING · IMPLEMENTATION", "데이터 통합부터 DINOv2 Tiny·배포까지 치아용으로 확장", 2)
cards = [
    ("01  DATA", "서로 다른 주석을 한 파이프라인으로", [
        ("입력", " Kaggle JSON / LabelMe / COCO"), ("변환", " polygon → box·mask·area"),
        ("안전성", " 빈 치아 이미지도 정상 학습")], CYAN),
    ("02  MODEL", "치아 번호 + 분할을 동시 예측", [
        ("클래스", " 32 teeth (+ background 처리)"), ("헤드", " box/class + segmentation head"),
        ("증강", " bbox/mask 동기 Albumentations")], TEAL),
    ("03  DEPLOY", "현업 추론 경로까지 연결", [
        ("복원", " resize 결과를 원본 크기로"), ("내보내기", " ONNX / CoreML 대응"),
        ("후처리", " threshold·NMS·FDI label")], YELLOW),
]
for i, (k, head, lines, col) in enumerate(cards):
    x = .6 + i*4.13
    rect(s, x, 1.52, 3.88, 3.38)
    text(s, k, x+.22, 1.75, 1.4, .25, 11, col, True)
    text(s, head, x+.22, 2.15, 3.42, .55, 17, WHITE, True)
    rich_lines(s, lines, x+.15, 2.89, 3.52, 1.65, 12.5, 8, True)

rect(s, .6, 5.13, 12.15, 1.80, RGBColor(11, 28, 47), LINE)
pill(s, "LIGHTWEIGHT", .86, 5.38, 1.65, TEAL, NAVY)
text(s, "DINOv2 Small → Tiny", 2.78, 5.32, 2.65, .35, 18, WHITE, True)
text(s, "hidden 384→192", 5.58, 5.35, 1.65, .28, 12.5, CYAN, True)
text(s, "heads 6→3", 7.28, 5.35, 1.25, .28, 12.5, CYAN, True)
text(s, "12 layers", 8.58, 5.35, 1.15, .28, 12.5, CYAN, True)
text(s, "Transformer 지배 연산/파라미터 ∝ 폭² → backbone 약 1/4 규모(구조 기준)", 2.78, 5.89, 7.25, .40, 12.5, MUTED)
codebox(s, "encoder = dino_tiny\npatch = 16\nresolution = 512\nqueries = 50", 10.20, 5.32, 2.30, 1.27, CYAN)
footer(s, "근거: develop...detr_teeth diff (+4,248/-41) · dinov2_tiny.json · cli/main.py · deploy/predict.py")

# 3 — reverse-registration architecture
s = prs.slides.add_slide(blank); set_bg(s); title(s, "REVERSE-REGISTRATION · ORIGINAL → CUSTOM", "원본 등록 코드에 X-ray ↔ IOS 전용 2D AI 정합 경로를 추가", 3)
rect(s, .6, 1.52, 3.25, 4.85)
text(s, "ORIGINAL", .86, 1.78, 1.1, .25, 11, MUTED, True)
text(s, "기초 3D/영상 등록", .86, 2.17, 2.55, .4, 20, WHITE, True)
rich_lines(s, [("입력", " CBCT / IOS / point cloud"), ("기반", " sampler·renderer·ICP"),
               ("한계", " 파노라마 X-ray 임상 경로 미완성")], .78, 2.84, 2.82, 1.75, 13, 10, True)
text(s, "origin/main", .86, 5.72, 1.35, .25, 11, MUTED)

arrow(s, 3.96, 3.36, .72, TEAL)
rect(s, 4.82, 1.52, 7.93, 4.85, RGBColor(12, 31, 50), LINE)
text(s, "CUSTOM · nonrigid2d", 5.08, 1.78, 2.3, .25, 11, TEAL, True)
text(s, "2단 AI 정합 + 임상 파이프라인", 5.08, 2.17, 4.6, .4, 20, WHITE, True)

steps = [("① COARSE", "SimilarityReg2D", "회전·이동·등방 scale\n+ arch (a,b) 추정", CYAN),
         ("② FINE", "NonRigidReg2D", "dense/anchor field\n+ 국소 치열 변형", TEAL),
         ("③ WORLD", "2D → 3D", "panoramic 좌표를\n+ world similarity로", YELLOW)]
for i, (k, name, body, col) in enumerate(steps):
    x = 5.08 + i*2.42
    rect(s, x, 2.92, 2.14, 1.75, RGBColor(7, 19, 33), LINE)
    text(s, k, x+.14, 3.10, 1.35, .22, 9.5, col, True)
    text(s, name, x+.14, 3.45, 1.87, .28, 15, WHITE, True)
    text(s, body, x+.14, 3.88, 1.84, .58, 11, MUTED)
    if i < 2: arrow(s, x+2.18, 3.60, .25, col)
rich_lines(s, [("데이터", " CocoImage2DDataset: point/polygon → EDT·signed EDT·mask"),
               ("학습", " supervised pose/deformation loss + iterative inference"),
               ("적용", " FDI 치아점 기반 상·하악 독립 정합 후 관계 합성")], 5.02, 4.94, 7.35, 1.16, 12.2, 4, True)
footer(s, "근거: origin/main...nonrigid2d (49 files, +17,062/-471) · models/similarity2d.py · models/deformation2d.py")

# 4 — reverse-registration engineering details
s = prs.slides.add_slide(blank); set_bg(s); title(s, "REVERSE-REGISTRATION · ENGINEERING", "치열 기하를 보존하고 실제 추론이 수렴하도록 안정화", 4)
features = [
    ("ARC-LENGTH", "치열 곡선 길이 보존", "x를 단순 affine하지 않고 곡선 길이를 기준으로 재매개화", CYAN),
    ("ANCHOR-AWARE", "치아 중심점 조건부 특징", "좌표 MLP / feature sampling + 순서 인지 attention pooling", TEAL),
    ("ARAP / ROBUST", "치아 강체성 보호", "shape loss + tooth-wise robust rigid 보정으로 과변형 억제", YELLOW),
    ("ITERATIVE", "반복 추론 제어", "deform은 1회, pose residual 위주로 반복해 오차 발산 방지", GREEN),
]
for i, (k, h, b, c) in enumerate(features):
    x = .6 + (i%2)*4.05; y = 1.52 + (i//2)*2.08
    rect(s, x, y, 3.78, 1.82)
    text(s, k, x+.2, y+.18, 1.5, .22, 9.5, c, True)
    text(s, h, x+.2, y+.52, 3.25, .3, 16, WHITE, True)
    text(s, b, x+.2, y+.96, 3.25, .68, 10.7, MUTED)

rect(s, 8.82, 1.52, 3.93, 4.43, RGBColor(7, 15, 27), LINE)
add_image_cover(s, REV / "result_draw.png", 9.02, 1.76, 3.53, 1.30, crop=False)
text(s, "IOS tooth anchors", 9.07, 3.15, 2.0, .25, 11, MUTED, True)
add_image_cover(s, REV / "upper.png", 9.02, 3.50, 3.53, 1.12, crop=False)
text(s, "X-ray/IOS → upper/lower → world transform", 9.07, 4.80, 3.15, .46, 12.5, WHITE, True)
pill(s, "CLINICAL INFERENCE", 9.07, 5.43, 2.02, TEAL, NAVY)

rect(s, .6, 5.88, 7.98, 1.03, RGBColor(14, 35, 56), LINE)
text(s, "검증 포인트", .86, 6.14, 1.05, .24, 11, TEAL, True)
text(s, "비단조 x 안전 처리 · pixel/denorm frame 보정 · signed-EDT 재구성 일치 · 단위 테스트 회귀 방지", 2.04, 6.10, 6.12, .35, 12.5, WHITE, True)
footer(s, "근거: pcregmodel/data/cococustom.py · models/losses.py · preproc/sampler.py · scripts/register_ai_process.py")

# 5 — RF-DETR core
s = prs.slides.add_slide(blank); set_bg(s); title(s, "RF-DETR · CORE TECHNOLOGY", "RF-DETR의 핵심: 강한 ViT 표현 + 희소한 end-to-end set prediction", 5)

flow = [("IMAGE", "입력/패딩", CYAN), ("DINOv2", "window attention", TEAL),
        ("PROJECTOR", "P3–P6 multi-scale", CYAN), ("DECODER", "MS deformable", YELLOW),
        ("SET OUTPUT", "class·box·mask", GREEN)]
for i, (a, b, c) in enumerate(flow):
    x = .62 + i*2.48
    rect(s, x, 1.56, 2.05, .92, RGBColor(9, 25, 42), c)
    text(s, a, x+.14, 1.75, 1.75, .23, 12, c, True, PP_ALIGN.CENTER)
    text(s, b, x+.10, 2.08, 1.83, .19, 9.5, MUTED, False, PP_ALIGN.CENTER)
    if i < 4: arrow(s, x+2.08, 1.82, .36, MUTED)

items = [
    ("1", "Windowed DINOv2", "local window 중심, 일부 블록만 global attention → 고해상도 비용 절감"),
    ("2", "ViTDet-style projector", "단일 ViT feature를 multi-scale memory로 변환"),
    ("3", "Two-stage query init", "proposal top-k 선택 후 decoder에서 box 반복 보정"),
    ("4", "Group-DETR training", "여러 query group을 독립 matching; 추론은 한 group"),
    ("5", "Set loss + optional mask", "분류·L1·GIoU·aux loss + query×pixel mask"),
    ("6", "Export-first path", "forward_export 분리로 ONNX/CoreML 추적 대응"),
]
for i, (n, h, b) in enumerate(items):
    x = .62 + (i%3)*4.12; y = 2.82 + (i//3)*1.65
    rect(s, x, y, 3.84, 1.39)
    pill(s, n, x+.18, y+.18, .36, TEAL, NAVY)
    text(s, h, x+.68, y+.18, 2.85, .28, 14.5, WHITE, True)
    text(s, b, x+.18, y+.62, 3.45, .56, 10.2, MUTED)

codebox(s, "backbone → transformer → class_embed / bbox_embed\nHungarianMatcher → SetCriterion → PostProcess", .62, 6.21, 7.65, .68, TEAL)
text(s, "핵심 차별점", 8.61, 6.27, 1.25, .22, 10, TEAL, True)
text(s, "ViT feature에서 바로 set prediction", 9.78, 6.24, 2.74, .40, 12.5, WHITE, True)
footer(s, "근거: models/backbone/* · models/transformer.py · models/lwdetr.py · models/matcher.py · segmentation_head.py")

prs.core_properties.title = "X-ray Teeth Numbering, Reverse Registration, RF-DETR 핵심 기술"
prs.core_properties.subject = "코드 기반 5장 기술 요약"
prs.core_properties.author = "Repository analysis"
prs.save(OUT / "xray_reverse_registration_rfdetr_기술요약.pptx")
print(OUT / "xray_reverse_registration_rfdetr_기술요약.pptx")
